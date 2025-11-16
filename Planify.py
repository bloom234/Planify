"""
Planify - Advanced AI Study Planner
"""

import streamlit as st
import pandas as pd
import json
import os
import io
import time
import random
import base64
from datetime import datetime, timedelta
from typing import Dict, List, Optional, Tuple

# Import required libraries
from dotenv import load_dotenv
import plotly.graph_objects as go
import plotly.express as px
from fpdf import FPDF
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import xlsxwriter
from PIL import Image, ImageDraw, ImageFont
import streamlit.components.v1 as components

# Load environment variables
load_dotenv()

# ==================== PAGE CONFIGURATION ====================
st.set_page_config(
    page_title="Planify - AI Study Planner",
    page_icon="🎯",
    layout="wide",
    initial_sidebar_state="expanded",
    menu_items={
        'Get Help': 'https://github.com/yourusername/planify',
        'About': "Planify - Your Personal AI Study Assistant"
    }
)

# ==================== CSS WITH ANIMATIONS ====================
def load_css():
    st.markdown("""
    <style>
    /* Import Google Fonts */
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700;800&family=Space+Grotesk:wght@400;700&display=swap');
    
    /* CSS Variables */
    :root {
        --primary: #6C63FF;
        --secondary: #FF6584;
        --success: #00BFA6;
        --warning: #FFA726;
        --danger: #EF5350;
        --dark: #2D3436;
        --light: #FFFFFF;
        --gray: #636E72;
        --bg-gradient: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
        --text-primary: #2D3436;
        --text-secondary: #636E72;
        --shadow-sm: 0 2px 4px rgba(0,0,0,0.06);
        --shadow-md: 0 4px 6px rgba(0,0,0,0.1);
        --shadow-lg: 0 10px 15px rgba(0,0,0,0.15);
        --shadow-xl: 0 20px 25px rgba(0,0,0,0.2);
        --animation-speed: 0.3s;
    }
    
    /* Global Styles */
    * {
        font-family: 'Inter', sans-serif;
        box-sizing: border-box;
    }
    
    /* Ensure text readability */
    .stApp {
        background: linear-gradient(135deg, #f5f7fa 0%, #c3cfe2 100%);
        color: var(--text-primary) !important;
    }
    
    p, span, div, label, li, td, th, input, textarea, select {
        color: var(--text-primary) !important;
    }
    
    h1, h2, h3, h4, h5, h6 {
        font-family: 'Space Grotesk', sans-serif !important;
        color: var(--text-primary) !important;
        font-weight: 700 !important;
    }
    
    /* Animated Hero Section */
    .hero-section {
        background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
        border-radius: 20px;
        padding: 3rem;
        margin-bottom: 2rem;
        position: relative;
        overflow: hidden;
        animation: slideDown 0.8s ease-out;
    }
    
    .hero-section::before {
        content: '';
        position: absolute;
        top: -50%;
        right: -50%;
        width: 200%;
        height: 200%;
        background: radial-gradient(circle, rgba(255,255,255,0.1) 0%, transparent 70%);
        animation: rotate 30s linear infinite;
    }
    
    .hero-section h1 {
        color: white !important;
        font-size: 3.5rem !important;
        margin: 0 !important;
        position: relative;
        z-index: 1;
        animation: fadeInUp 0.8s ease-out 0.2s both;
    }
    
    .hero-section p {
        color: rgba(255,255,255,0.9) !important;
        font-size: 1.2rem !important;
        margin-top: 1rem !important;
        position: relative;
        z-index: 1;
        animation: fadeInUp 0.8s ease-out 0.4s both;
    }
    
    /* Animations */
    @keyframes slideDown {
        from {
            transform: translateY(-100px);
            opacity: 0;
        }
        to {
            transform: translateY(0);
            opacity: 1;
        }
    }
    
    @keyframes fadeInUp {
        from {
            transform: translateY(30px);
            opacity: 0;
        }
        to {
            transform: translateY(0);
            opacity: 1;
        }
    }
    
    @keyframes rotate {
        from { transform: rotate(0deg); }
        to { transform: rotate(360deg); }
    }
    
    @keyframes pulse {
        0%, 100% { transform: scale(1); }
        50% { transform: scale(1.05); }
    }
    
    @keyframes bounce {
        0%, 100% { transform: translateY(0); }
        50% { transform: translateY(-10px); }
    }
    
    /* Custom Loader */
    .loader-wrapper {
        display: flex;
        justify-content: center;
        align-items: center;
        height: 200px;
    }
    
    .custom-loader {
        width: 50px;
        height: 50px;
        position: relative;
    }
    
    .custom-loader div {
        position: absolute;
        width: 100%;
        height: 100%;
        border-radius: 50%;
        background: var(--primary);
        opacity: 0.6;
        animation: ripple 1.5s infinite;
    }
    
    .custom-loader div:nth-child(2) {
        animation-delay: -0.5s;
    }
    
    @keyframes ripple {
        0% {
            transform: scale(0);
            opacity: 1;
        }
        100% {
            transform: scale(1.5);
            opacity: 0;
        }
    }
    
    /* Chat Messages */
    .chat-container {
        max-width: 900px;
        margin: 0 auto;
    }
    
    .chat-message {
        margin: 1.5rem 0;
        animation: messageSlide 0.5s ease-out;
    }
    
    @keyframes messageSlide {
        from {
            opacity: 0;
            transform: translateX(-20px);
        }
        to {
            opacity: 1;
            transform: translateX(0);
        }
    }
    
    .bot-message {
        background: white;
        border-left: 4px solid var(--primary);
        padding: 1.5rem;
        border-radius: 0 15px 15px 15px;
        box-shadow: var(--shadow-md);
        color: var(--text-primary) !important;
        position: relative;
        margin-right: 15%;
    }
    
    .bot-message::before {
        content: '🤖';
        position: absolute;
        left: -40px;
        top: 20px;
        font-size: 24px;
        animation: bounce 2s infinite;
    }
    
    .user-message {
        background: linear-gradient(135deg, var(--primary) 0%, #9D50BB 100%);
        color: white !important;
        padding: 1.5rem;
        border-radius: 15px 0 15px 15px;
        box-shadow: var(--shadow-md);
        margin-left: 15%;
        position: relative;
    }
    
    .user-message * {
        color: white !important;
    }
    
    .user-message::after {
        content: '👤';
        position: absolute;
        right: -40px;
        top: 20px;
        font-size: 24px;
    }
    
    /* Progress Steps */
    .progress-wrapper {
        background: white;
        border-radius: 15px;
        padding: 2rem;
        margin: 2rem 0;
        box-shadow: var(--shadow-lg);
    }
    
    .progress-step {
        display: inline-block;
        padding: 0.75rem 1.5rem;
        margin: 0.25rem;
        border-radius: 25px;
        background: #f0f0f0;
        color: var(--text-primary);
        font-weight: 600;
        transition: all 0.3s;
    }
    
    .progress-step.active {
        background: var(--primary);
        color: white !important;
        animation: pulse 1.5s infinite;
        box-shadow: 0 0 20px rgba(108, 99, 255, 0.4);
    }
    
    .progress-step.completed {
        background: var(--success);
        color: white !important;
    }
    
    /* Template Cards */
    .template-card {
        padding: 2rem;
        border-radius: 20px;
        transition: all 0.3s;
        cursor: pointer;
        position: relative;
        overflow: hidden;
        margin: 1rem 0;
    }
    
    .template-simple {
        background: white;
        border: 2px solid #e0e0e0;
    }
    
    .template-minimal {
        background: linear-gradient(135deg, #f5f7fa 0%, #c3cfe2 100%);
        border: none;
        box-shadow: var(--shadow-md);
    }
    
    .template-aesthetic {
        background: linear-gradient(135deg, #FA8BFF 0%, #2BD2FF 52%, #2BFF88 90%);
        border: none;
        box-shadow: var(--shadow-lg);
        color: white;
    }
    
    .template-card:hover {
        transform: translateY(-5px) scale(1.02);
        box-shadow: var(--shadow-xl);
    }
    
    /* Buttons */
    .stButton > button {
        background: linear-gradient(135deg, var(--primary) 0%, #9D50BB 100%);
        color: white !important;
        border: none;
        padding: 0.75rem 2rem;
        border-radius: 30px;
        font-weight: 600;
        transition: all 0.3s;
        box-shadow: var(--shadow-md);
        position: relative;
        overflow: hidden;
    }
    
    .stButton > button::before {
        content: '';
        position: absolute;
        top: 0;
        left: -100%;
        width: 100%;
        height: 100%;
        background: rgba(255,255,255,0.3);
        transition: left 0.5s;
    }
    
    .stButton > button:hover {
        transform: translateY(-2px);
        box-shadow: var(--shadow-xl);
    }
    
    .stButton > button:hover::before {
        left: 100%;
    }
    
    /* Input Fields */
    .stTextInput input,
    .stTextArea textarea,
    .stSelectbox select,
    .stTimeInput input {
        background: white !important;
        color: var(--text-primary) !important;
        border: 2px solid #e0e0e0 !important;
        border-radius: 10px !important;
        padding: 0.75rem !important;
        font-size: 14px !important;
        transition: all 0.3s !important;
    }
    
    .stTextInput input:focus,
    .stTextArea textarea:focus {
        border-color: var(--primary) !important;
        box-shadow: 0 0 0 3px rgba(108, 99, 255, 0.1) !important;
        outline: none !important;
    }
    
    /* Download Buttons */
    .stDownloadButton > button {
        background: linear-gradient(135deg, var(--success) 0%, #00E676 100%);
        color: white !important;
        border: none;
        padding: 0.75rem 1.5rem;
        border-radius: 10px;
        font-weight: 600;
        box-shadow: var(--shadow-md);
    }
    
    .stDownloadButton > button:hover {
        transform: translateY(-2px);
        box-shadow: var(--shadow-lg);
    }
    
    /* Success Animation */
    .success-message {
        text-align: center;
        padding: 2rem;
        animation: zoomIn 0.5s ease-out;
    }
    
    @keyframes zoomIn {
        from {
            transform: scale(0);
            opacity: 0;
        }
        to {
            transform: scale(1);
            opacity: 1;
        }
    }
    
    /* Floating Button */
    .float-button {
        position: fixed;
        bottom: 30px;
        right: 30px;
        width: 60px;
        height: 60px;
        background: var(--primary);
        border-radius: 50%;
        display: flex;
        align-items: center;
        justify-content: center;
        color: white;
        font-size: 24px;
        box-shadow: var(--shadow-lg);
        cursor: pointer;
        z-index: 1000;
        animation: floatButton 2s ease-in-out infinite;
    }
    
    @keyframes floatButton {
        0%, 100% { transform: translateY(0); }
        50% { transform: translateY(-10px); }
    }
    
    /* Responsive Design */
    @media (max-width: 768px) {
        .hero-section h1 {
            font-size: 2.5rem !important;
        }
        
        .chat-message {
            margin-left: 0 !important;
            margin-right: 0 !important;
        }
    }
    </style>
    """, unsafe_allow_html=True)

# ==================== LOTTIE ANIMATIONS ====================
def load_lottie_animation():
    lottie_code = """
    <script src="https://unpkg.com/@lottiefiles/lottie-player@latest/dist/lottie-player.js"></script>
    <lottie-player 
        src="https://assets2.lottiefiles.com/packages/lf20_V9t630.json"
        background="transparent"
        speed="1"
        style="width: 300px; height: 300px; margin: 0 auto;"
        loop
        autoplay>
    </lottie-player>
    """
    return lottie_code

# ==================== AI PROVIDER CLASS ====================
class AIProvider:
    """Unified AI Provider for Groq and OpenAI"""
    
    def __init__(self):
        self.provider = None
        self.client = None
        self._initialize()
    
    def _initialize(self):
        """Initialize available AI provider"""
        # Try Groq first
        if os.getenv('GROQ_API_KEY'):
            try:
                from groq import Groq
                self.client = Groq(api_key=os.getenv('GROQ_API_KEY'))
                self.provider = 'groq'
                st.success("✅ Connected to Groq AI")
            except Exception as e:
                st.warning(f"Groq initialization failed: {e}")
        
        # Try OpenAI if Groq not available
        if not self.provider and os.getenv('OPENAI_API_KEY'):
            try:
                import openai
                openai.api_key = os.getenv('OPENAI_API_KEY')
                self.client = openai
                self.provider = 'openai'
                st.success("✅ Connected to OpenAI")
            except Exception as e:
                st.warning(f"OpenAI initialization failed: {e}")
        
        # Fallback mode
        if not self.provider:
            self.provider = 'offline'
            st.info("ℹ️ Running in offline mode")
    
    def chat(self, prompt: str, context: List = None) -> str:
        """Get AI response"""
        if self.provider == 'groq':
            return self._groq_chat(prompt, context)
        elif self.provider == 'openai':
            return self._openai_chat(prompt, context)
        else:
            return self._offline_response(prompt)
    
    def _groq_chat(self, prompt: str, context: List) -> str:
        """Chat using Groq"""
        try:
            messages = [
                {"role": "system", "content": "You are Planify, a friendly AI study planner assistant. Help students create personalized study schedules."}
            ]
            if context:
                messages.extend(context[-5:])  # Keep last 5 messages
            messages.append({"role": "user", "content": prompt})
            
            response = self.client.chat.completions.create(
                model="llama-3.3-70b-versatile",
                messages=messages,
                temperature=0.7,
                max_tokens=1000
            )
            return response.choices[0].message.content
        except:
            return self._offline_response(prompt)
    
    def _openai_chat(self, prompt: str, context: List) -> str:
        """Chat using OpenAI"""
        try:
            import openai
            messages = [
                {"role": "system", "content": "You are Planify, a friendly AI study planner assistant."}
            ]
            if context:
                messages.extend(context[-5:])
            messages.append({"role": "user", "content": prompt})
            
            response = openai.ChatCompletion.create(
                model="gpt-3.5-turbo",
                messages=messages,
                temperature=0.7,
                max_tokens=500
            )
            return response['choices'][0]['message']['content']
        except:
            return self._offline_response(prompt)
    
    def _offline_response(self, prompt: str) -> str:
        """Offline fallback responses"""
        responses = {
            "greeting": "Hello! I'm Planify, your AI study assistant. Let's create your perfect study plan!",
            "name": "Great choice! Let's move forward with your plan.",
            "type": "Excellent selection! This will help structure your studies effectively.",
            "problem": "I understand your challenge. We'll address this in your custom plan.",
            "routine": "Thanks for sharing your routine. I'll optimize your schedule accordingly.",
            "subjects": "Perfect! I'll organize these subjects for maximum efficiency.",
            "template": "Beautiful choice! Your planner will look amazing.",
            "success": "Your personalized planner is ready!"
        }
        
        for key in responses:
            if key in prompt.lower():
                return responses[key]
        
        return "Let's continue building your perfect study plan!"

# ==================== SCHEDULE GENERATOR ====================
class ScheduleGenerator:
    """Generate customized study schedules"""
    
    @staticmethod
    def create_schedule(data: Dict) -> pd.DataFrame:
        """Create schedule based on user data"""
        plan_type = data.get('plan_type', 'daily')
        
        if plan_type == 'daily':
            return ScheduleGenerator._daily_schedule(data)
        elif plan_type == 'weekly':
            return ScheduleGenerator._weekly_schedule(data)
        else:
            return ScheduleGenerator._monthly_schedule(data)
    
    @staticmethod
    def _daily_schedule(data: Dict) -> pd.DataFrame:
        """Generate daily schedule"""
        routine = data.get('routine', {})
        subjects = data.get('subjects', ['Math', 'Science', 'English'])
        
        schedule = []
        
        # Morning routine
        wake_time = routine.get('wake_time', '07:00')
        schedule.append({
            'Time': wake_time,
            'Activity': '🌅 Wake Up & Morning Routine',
            'Duration': '30 min',
            'Type': 'Personal',
            'Energy Level': '🔋 Building'
        })
        
        # Breakfast
        breakfast_time = routine.get('breakfast_time', '08:00')
        schedule.append({
            'Time': breakfast_time,
            'Activity': '🍳 Breakfast',
            'Duration': '30 min',
            'Type': 'Meal',
            'Energy Level': '🔋🔋 Good'
        })
        
        # Study sessions based on user preferences
        study_sessions = routine.get('study_sessions', [])
        if study_sessions:
            for i, session in enumerate(study_sessions):
                subject = subjects[i % len(subjects)] if subjects else f'Subject {i+1}'
                schedule.append({
                    'Time': session.get('start_time', f'{9+i*2}:00'),
                    'Activity': f'📚 Study: {subject}',
                    'Duration': session.get('duration', '2 hours'),
                    'Type': 'Study',
                    'Energy Level': '🔋🔋🔋 Peak'
                })
                
                # Add break
                schedule.append({
                    'Time': session.get('break_time', f'{10+i*2}:45'),
                    'Activity': '☕ Break',
                    'Duration': '15 min',
                    'Type': 'Break',
                    'Energy Level': '🔋 Recharge'
                })
        else:
            # Default study sessions
            times = ['09:00', '11:30', '14:00', '16:30', '19:00']
            for i, time in enumerate(times[:len(subjects)]):
                schedule.append({
                    'Time': time,
                    'Activity': f'📚 Study: {subjects[i % len(subjects)]}',
                    'Duration': '2 hours',
                    'Type': 'Study',
                    'Energy Level': '🔋🔋🔋 Peak'
                })
        
        # Lunch
        lunch_time = routine.get('lunch_time', '13:00')
        schedule.append({
            'Time': lunch_time,
            'Activity': '🍽️ Lunch',
            'Duration': '45 min',
            'Type': 'Meal',
            'Energy Level': '🔋🔋 Good'
        })
        
        # Dinner
        dinner_time = routine.get('dinner_time', '19:30')
        schedule.append({
            'Time': dinner_time,
            'Activity': '🍝 Dinner',
            'Duration': '45 min',
            'Type': 'Meal',
            'Energy Level': '🔋🔋 Good'
        })
        
        # Sleep
        sleep_time = routine.get('sleep_time', '22:30')
        schedule.append({
            'Time': sleep_time,
            'Activity': '😴 Sleep Preparation',
            'Duration': '30 min',
            'Type': 'Personal',
            'Energy Level': '🔋 Winding Down'
        })
        
        # Sort by time
        df = pd.DataFrame(schedule)
        df['Time'] = pd.to_datetime(df['Time'], format='%H:%M').dt.time
        df = df.sort_values('Time')
        df['Time'] = df['Time'].astype(str).str[:5]
        
        return df
    
    @staticmethod
    def _weekly_schedule(data: Dict) -> pd.DataFrame:
        """Generate weekly schedule"""
        subjects = data.get('subjects', ['Math', 'Science', 'English'])
        
        days = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday', 'Sunday']
        schedule = []
        
        for i, day in enumerate(days):
            if day == 'Sunday':
                schedule.append({
                    'Day': day,
                    'Morning (7-12)': '🌅 Rest/Light Review',
                    'Afternoon (12-5)': '📝 Weekly Planning',
                    'Evening (5-10)': '🎯 Goal Setting',
                    'Focus Subject': 'Review Week',
                    'Special Notes': '✨ Recharge Day'
                })
            else:
                subject_index = i % len(subjects)
                schedule.append({
                    'Day': day,
                    'Morning (7-12)': f'📚 {subjects[subject_index]}',
                    'Afternoon (12-5)': f'✍️ Practice {subjects[(subject_index + 1) % len(subjects)]}',
                    'Evening (5-10)': '📖 Review & Homework',
                    'Focus Subject': subjects[subject_index],
                    'Special Notes': '💪 Stay Focused!'
                })
        
        return pd.DataFrame(schedule)
    
    @staticmethod
    def _monthly_schedule(data: Dict) -> pd.DataFrame:
        """Generate monthly schedule"""
        subjects = data.get('subjects', ['Math', 'Science', 'English'])
        
        weeks = ['Week 1', 'Week 2', 'Week 3', 'Week 4']
        schedule = []
        
        for i, week in enumerate(weeks):
            schedule.append({
                'Week': week,
                'Phase': ['Foundation', 'Development', 'Advanced', 'Revision'][i],
                'Focus Areas': ', '.join(subjects),
                'Target': f'{(i+1)*25}% Complete',
                'Milestones': ['Basics Clear', 'Practice Started', 'Mock Tests', 'Final Review'][i],
                'Motivation': ['🚀 Strong Start!', '💪 Keep Going!', '🎯 Almost There!', '🏆 Final Push!'][i]
            })
        
        return pd.DataFrame(schedule)

# ==================== TEMPLATE STYLER ====================
class TemplateStyler:
    """Apply different visual styles to schedules"""
    
    @staticmethod
    def apply_style(df: pd.DataFrame, template: str) -> pd.DataFrame:
        """Apply template styling"""
        if template == 'aesthetic':
            return TemplateStyler._aesthetic_style(df)
        elif template == 'minimal':
            return TemplateStyler._minimal_style(df)
        else:
            return df  # Simple style - no modifications
    
    @staticmethod
    def _aesthetic_style(df: pd.DataFrame) -> pd.DataFrame:
        """Apply aesthetic styling with emojis and colors"""
        styled_df = df.copy()
        
        # Add decorative elements
        emoji_decorations = ['✨', '🌟', '💫', '⭐', '🌈']
        
        # Add random decorative emojis to some cells
        for col in styled_df.columns:
            if col not in ['Time', 'Day', 'Week']:
                for idx in styled_df.index:
                    if random.random() > 0.7:  # 30% chance
                        current_val = str(styled_df.at[idx, col])
                        decoration = random.choice(emoji_decorations)
                        styled_df.at[idx, col] = f"{current_val} {decoration}"
        
        return styled_df
    
    @staticmethod
    def _minimal_style(df: pd.DataFrame) -> pd.DataFrame:
        """Apply minimal clean styling"""
        styled_df = df.copy()
        
        # Remove emojis for minimal look
        for col in styled_df.columns:
            styled_df[col] = styled_df[col].astype(str).str.replace(r'[^\w\s:]', '', regex=True)
        
        return styled_df

# ==================== EXPORT MANAGER ====================
class ExportManager:
    """Handle all export operations"""
    
    @staticmethod
    def to_pdf(df: pd.DataFrame, data: Dict) -> bytes:
        """Export to PDF"""
        try:
            pdf = FPDF()
            pdf.add_page()
            pdf.set_auto_page_break(auto=True, margin=15)
            
            # Add custom font if available
            pdf.set_font("Arial", size=12)
            
            # Title
            pdf.set_font("Arial", 'B', 24)
            pdf.cell(200, 10, txt="Planify Study Planner", ln=True, align='C')
            
            # Subtitle
            pdf.set_font("Arial", 'I', 14)
            template_name = data.get('template', 'Simple').title()
            plan_type = data.get('plan_type', 'Daily').title()
            pdf.cell(200, 10, txt=f"{plan_type} Schedule - {template_name} Style", ln=True, align='C')
            
            # Add space
            pdf.ln(10)
            
            # Project info
            pdf.set_font("Arial", size=11)
            pdf.cell(200, 10, txt=f"Project: {data.get('folder_name', 'My Plan')}", ln=True)
            pdf.cell(200, 10, txt=f"Created: {datetime.now().strftime('%B %d, %Y')}", ln=True)
            
            # Add space before table
            pdf.ln(10)
            
            # Calculate column widths
            page_width = pdf.w - 2 * pdf.l_margin
            col_count = len(df.columns)
            col_width = page_width / col_count
            
            # Table header
            pdf.set_font("Arial", 'B', 10)
            pdf.set_fill_color(108, 99, 255)  # Primary color
            pdf.set_text_color(255, 255, 255)
            
            for col in df.columns:
                pdf.cell(col_width, 10, str(col), 1, 0, 'C', True)
            pdf.ln()
            
            # Table data
            pdf.set_font("Arial", size=9)
            pdf.set_text_color(0, 0, 0)
            
            for _, row in df.iterrows():
                for col in df.columns:
                    value = str(row[col])
                    # Truncate long text
                    if len(value) > 20:
                        value = value[:17] + "..."
                    pdf.cell(col_width, 8, value, 1, 0, 'C')
                pdf.ln()
            
            # Add motivational quote
            pdf.ln(10)
            pdf.set_font("Arial", 'I', 11)
            quotes = [
                "Success is the sum of small efforts repeated day in and day out.",
                "The expert in anything was once a beginner.",
                "Focus on progress, not perfection."
            ]
            pdf.multi_cell(0, 10, random.choice(quotes), align='C')
            
            return pdf.output(dest='S').encode('latin-1')
            
        except Exception as e:
            st.error(f"PDF generation error: {e}")
            return None
    
    @staticmethod
    def to_excel(df: pd.DataFrame, data: Dict) -> bytes:
        """Export to Excel with styling"""
        try:
            output = io.BytesIO()
            
            with pd.ExcelWriter(output, engine='xlsxwriter') as writer:
                # Write main sheet
                df.to_excel(writer, sheet_name='Schedule', index=False)
                
                # Get workbook and worksheet
                workbook = writer.book
                worksheet = writer.sheets['Schedule']
                
                # Define formats based on template
                template = data.get('template', 'simple')
                
                if template == 'aesthetic':
                    # Aesthetic format
                    header_format = workbook.add_format({
                        'bold': True,
                        'text_wrap': True,
                        'valign': 'center',
                        'align': 'center',
                        'fg_color': '#FF6B9D',
                        'font_color': 'white',
                        'border': 1,
                        'font_size': 12
                    })
                    
                    cell_format = workbook.add_format({
                        'text_wrap': True,
                        'valign': 'center',
                        'align': 'center',
                        'border': 1,
                        'fg_color': '#FFF0F5'
                    })
                    
                    # Apply gradient-like effect with alternating colors
                    alt_format = workbook.add_format({
                        'text_wrap': True,
                        'valign': 'center',
                        'align': 'center',
                        'border': 1,
                        'fg_color': '#FFE0EC'
                    })
                    
                elif template == 'minimal':
                    # Minimal format
                    header_format = workbook.add_format({
                        'bold': True,
                        'valign': 'center',
                        'align': 'center',
                        'fg_color': '#F5F5F5',
                        'border': 1
                    })
                    
                    cell_format = workbook.add_format({
                        'valign': 'center',
                        'align': 'center',
                        'border': 1
                    })
                    
                    alt_format = cell_format
                    
                else:
                    # Simple format
                    header_format = workbook.add_format({
                        'bold': True,
                        'border': 1
                    })
                    
                    cell_format = workbook.add_format({
                        'border': 1
                    })
                    
                    alt_format = cell_format
                
                # Write headers
                for col_num, value in enumerate(df.columns):
                    worksheet.write(0, col_num, value, header_format)
                
                # Write data with alternating row colors for aesthetic
                for row_num, row_data in enumerate(df.values):
                    format_to_use = alt_format if (template == 'aesthetic' and row_num % 2 == 0) else cell_format
                    for col_num, value in enumerate(row_data):
                        worksheet.write(row_num + 1, col_num, str(value), format_to_use)
                
                # Adjust column widths
                for i, col in enumerate(df.columns):
                    max_length = max(df[col].astype(str).map(len).max(), len(col))
                    worksheet.set_column(i, i, min(max_length + 2, 30))
                
                # Add project info sheet
                info_df = pd.DataFrame({
                    'Property': ['Project Name', 'Type', 'Template', 'Created'],
                    'Value': [
                        data.get('folder_name', 'My Plan'),
                        data.get('plan_type', 'daily'),
                        data.get('template', 'simple'),
                        datetime.now().strftime('%Y-%m-%d %H:%M')
                    ]
                })
                info_df.to_excel(writer, sheet_name='Info', index=False)
            
            output.seek(0)
            return output.getvalue()
            
        except Exception as e:
            st.error(f"Excel generation error: {e}")
            return None
    
    @staticmethod
    def to_csv(df: pd.DataFrame) -> bytes:
        """Export to CSV"""
        try:
            return df.to_csv(index=False).encode('utf-8')
        except Exception as e:
            st.error(f"CSV generation error: {e}")
            return None

# ==================== UI COMPONENTS ====================
def show_loader(message: str = "Processing...", duration: float = 2):
    """Display animated loader"""
    placeholder = st.empty()
    with placeholder.container():
        st.markdown(f"""
        <div class="loader-wrapper">
            <div class="custom-loader">
                <div></div>
                <div></div>
            </div>
        </div>
        <h3 style="text-align: center; color: var(--text-primary); margin-top: 1rem;">
            {message}
        </h3>
        """, unsafe_allow_html=True)
    time.sleep(duration)
    placeholder.empty()

def show_hero_section():
    """Display hero section"""
    st.markdown("""
    <div class="hero-section">
        <h1>🎯 Planify</h1>
        <p>Your Personal AI-Powered Study Planning Assistant</p>
        <p style="font-size: 1rem; opacity: 0.9;">Create smart schedules that adapt to your life</p>
    </div>
    """, unsafe_allow_html=True)

def show_progress(step: int):
    """Display progress steps"""
    steps = [
        ("📁", "Project"),
        ("📅", "Type"),
        ("💭", "Challenge"),
        ("⏰", "Routine"),
        ("📚", "Subjects"),
        ("🎨", "Style"),
        ("✨", "Generate")
    ]
    
    st.markdown('<div class="progress-wrapper">', unsafe_allow_html=True)
    cols = st.columns(len(steps))
    
    for i, (col, (icon, label)) in enumerate(zip(cols, steps)):
        step_num = i + 1
        with col:
            if step_num < step:
                status = "completed"
                display_icon = "✅"
            elif step_num == step:
                status = "active"
                display_icon = icon
            else:
                status = ""
                display_icon = icon
            
            st.markdown(f"""
            <div class="progress-step {status}">
                {display_icon} {label}
            </div>
            """, unsafe_allow_html=True)
    
    st.markdown('</div>', unsafe_allow_html=True)

def show_success_message(message: str):
    """Display success animation"""
    st.markdown(f"""
    <div class="success-message">
        <h1 style="font-size: 72px; margin: 0;">🎉</h1>
        <h2 style="color: var(--text-primary); margin-top: 1rem;">{message}</h2>
        <p style="color: var(--text-secondary); margin-top: 0.5rem;">Your personalized study plan is ready!</p>
    </div>
    """, unsafe_allow_html=True)

# ==================== SESSION STATE ====================
def init_session_state():
    """Initialize session state variables"""
    if 'step' not in st.session_state:
        st.session_state.step = 1
    if 'messages' not in st.session_state:
        st.session_state.messages = []
    if 'project_data' not in st.session_state:
        st.session_state.project_data = {
            'folder_name': '',
            'plan_type': '',
            'problem': '',
            'routine': {},
            'subjects': [],
            'template': '',
            'generated_plan': None
        }
    if 'ai_provider' not in st.session_state:
        st.session_state.ai_provider = AIProvider()
    if 'conversation_context' not in st.session_state:
        st.session_state.conversation_context = []

# ==================== MAIN APPLICATION ====================
def main():
    # Load CSS
    load_css()
    
    # Initialize session state
    init_session_state()
    
    # Hero section
    show_hero_section()
    
    # Progress indicator
    show_progress(st.session_state.step)
    
    # Main container
    container = st.container()
    
    with container:
        # Display chat messages
        for msg in st.session_state.messages:
            if msg["role"] == "assistant":
                st.markdown(f'<div class="chat-message bot-message">{msg["content"]}</div>', 
                          unsafe_allow_html=True)
            else:
                st.markdown(f'<div class="chat-message user-message">{msg["content"]}</div>', 
                          unsafe_allow_html=True)
        
        # Step 1: Project Name
        if st.session_state.step == 1:
            if not st.session_state.messages:
                welcome = "👋 Hello! I'm Planify, your AI study assistant. Let's create your perfect study plan! What would you like to name your project?"
                st.session_state.messages.append({"role": "assistant", "content": welcome})
                st.rerun()
            
            col1, col2 = st.columns([3, 1])
            with col1:
                project_name = st.text_input(
                    "Project Name",
                    placeholder="e.g., 'Final Exam Preparation', 'Weekly Study Plan'",
                    key="project_name_input"
                )
            with col2:
                st.write("")  # Spacer
                st.write("")  # Spacer
                if st.button("Continue →", key="btn1", use_container_width=True):
                    if project_name:
                        st.session_state.project_data['folder_name'] = project_name
                        st.session_state.messages.append({"role": "user", "content": project_name})
                        response = f"Great! I've created '{project_name}' for you. Now, what type of planner would work best for your needs?"
                        st.session_state.messages.append({"role": "assistant", "content": response})
                        st.session_state.step = 2
                        st.rerun()
                    else:
                        st.error("Please enter a project name")
        
        # Step 2: Plan Type
        elif st.session_state.step == 2:
            st.markdown("### Choose Your Planning Style")
            
            col1, col2, col3 = st.columns(3)
            
            with col1:
                st.markdown("""
                <div style="text-align: center; padding: 1rem;">
                    <h1 style="font-size: 48px; margin: 0;">📅</h1>
                    <h4>Daily Planner</h4>
                    <p style="color: var(--text-secondary); font-size: 14px;">
                        Hour-by-hour schedule for maximum productivity
                    </p>
                </div>
                """, unsafe_allow_html=True)
                if st.button("Choose Daily", key="daily", use_container_width=True):
                    st.session_state.project_data['plan_type'] = 'daily'
                    st.session_state.messages.append({"role": "user", "content": "Daily planner"})
                    response = "Perfect! A daily planner will help you manage every hour effectively. What challenges do you face while studying?"
                    st.session_state.messages.append({"role": "assistant", "content": response})
                    st.session_state.step = 3
                    st.rerun()
            
            with col2:
                st.markdown("""
                <div style="text-align: center; padding: 1rem;">
                    <h1 style="font-size: 48px; margin: 0;">📆</h1>
                    <h4>Weekly Planner</h4>
                    <p style="color: var(--text-secondary); font-size: 14px;">
                        7-day overview for balanced learning
                    </p>
                </div>
                """, unsafe_allow_html=True)
                if st.button("Choose Weekly", key="weekly", use_container_width=True):
                    st.session_state.project_data['plan_type'] = 'weekly'
                    st.session_state.messages.append({"role": "user", "content": "Weekly planner"})
                    response = "Excellent! A weekly planner provides great balance and flexibility. What study challenges should we address?"
                    st.session_state.messages.append({"role": "assistant", "content": response})
                    st.session_state.step = 3
                    st.rerun()
            
            with col3:
                st.markdown("""
                <div style="text-align: center; padding: 1rem;">
                    <h1 style="font-size: 48px; margin: 0;">🗓️</h1>
                    <h4>Monthly Planner</h4>
                    <p style="color: var(--text-secondary); font-size: 14px;">
                        Long-term goals and milestones
                    </p>
                </div>
                """, unsafe_allow_html=True)
                if st.button("Choose Monthly", key="monthly", use_container_width=True):
                    st.session_state.project_data['plan_type'] = 'monthly'
                    st.session_state.messages.append({"role": "user", "content": "Monthly planner"})
                    response = "Great choice! Monthly planning helps track long-term progress. What challenges do you face in your studies?"
                    st.session_state.messages.append({"role": "assistant", "content": response})
                    st.session_state.step = 3
                    st.rerun()
        
        # Step 3: Problem/Challenge
        elif st.session_state.step == 3:
            st.markdown("### Share Your Challenges")
            
            problem = st.text_area(
                "What difficulties do you face while studying?",
                placeholder="e.g., 'I get distracted easily', 'Hard to manage multiple subjects', 'Procrastination issues'",
                height=100,
                key="problem_input"
            )
            
            if st.button("Continue →", key="btn3", use_container_width=True):
                if problem:
                    st.session_state.project_data['problem'] = problem
                    st.session_state.messages.append({"role": "user", "content": problem})
                    response = "I understand your challenges. I'll make sure your plan addresses these issues. Now, let's talk about your daily routine. What time do you usually wake up?"
                    st.session_state.messages.append({"role": "assistant", "content": response})
                    st.session_state.step = 4
                    st.rerun()
                else:
                    st.error("Please describe your challenges")
        
        # Step 4: Routine (Conversational)
        elif st.session_state.step == 4:
            st.markdown("### Let's Build Your Routine")
            
            # Collect routine information step by step
            if 'routine_step' not in st.session_state:
                st.session_state.routine_step = 'wake'
            
            if st.session_state.routine_step == 'wake':
                wake_time = st.time_input("What time do you wake up?", 
                                         value=datetime.strptime("07:00", "%H:%M").time())
                if st.button("Next →", key="wake_btn"):
                    st.session_state.project_data['routine']['wake_time'] = wake_time.strftime("%H:%M")
                    st.session_state.messages.append({"role": "user", 
                                                     "content": f"I wake up at {wake_time.strftime('%H:%M')}"})
                    st.session_state.messages.append({"role": "assistant", 
                                                     "content": "Good! When do you usually have breakfast?"})
                    st.session_state.routine_step = 'breakfast'
                    st.rerun()
            
            elif st.session_state.routine_step == 'breakfast':
                breakfast_time = st.time_input("Breakfast time:", 
                                              value=datetime.strptime("08:00", "%H:%M").time())
                if st.button("Next →", key="breakfast_btn"):
                    st.session_state.project_data['routine']['breakfast_time'] = breakfast_time.strftime("%H:%M")
                    st.session_state.messages.append({"role": "user", 
                                                     "content": f"Breakfast at {breakfast_time.strftime('%H:%M')}"})
                    st.session_state.messages.append({"role": "assistant", 
                                                     "content": "When do you prefer to have lunch?"})
                    st.session_state.routine_step = 'lunch'
                    st.rerun()
            
            elif st.session_state.routine_step == 'lunch':
                lunch_time = st.time_input("Lunch time:", 
                                          value=datetime.strptime("13:00", "%H:%M").time())
                if st.button("Next →", key="lunch_btn"):
                    st.session_state.project_data['routine']['lunch_time'] = lunch_time.strftime("%H:%M")
                    st.session_state.messages.append({"role": "user", 
                                                     "content": f"Lunch at {lunch_time.strftime('%H:%M')}"})
                    st.session_state.messages.append({"role": "assistant", 
                                                     "content": "What time is dinner?"})
                    st.session_state.routine_step = 'dinner'
                    st.rerun()
            
            elif st.session_state.routine_step == 'dinner':
                dinner_time = st.time_input("Dinner time:", 
                                           value=datetime.strptime("19:30", "%H:%M").time())
                if st.button("Next →", key="dinner_btn"):
                    st.session_state.project_data['routine']['dinner_time'] = dinner_time.strftime("%H:%M")
                    st.session_state.messages.append({"role": "user", 
                                                     "content": f"Dinner at {dinner_time.strftime('%H:%M')}"})
                    st.session_state.messages.append({"role": "assistant", 
                                                     "content": "What time do you go to sleep?"})
                    st.session_state.routine_step = 'sleep'
                    st.rerun()
            
            elif st.session_state.routine_step == 'sleep':
                sleep_time = st.time_input("Sleep time:", 
                                          value=datetime.strptime("22:30", "%H:%M").time())
                if st.button("Next →", key="sleep_btn"):
                    st.session_state.project_data['routine']['sleep_time'] = sleep_time.strftime("%H:%M")
                    st.session_state.messages.append({"role": "user", 
                                                     "content": f"I sleep at {sleep_time.strftime('%H:%M')}"})
                    st.session_state.messages.append({"role": "assistant", 
                                                     "content": "Perfect! I have your daily routine. Now, when do you prefer to study? Morning, afternoon, or evening?"})
                    st.session_state.routine_step = 'study_pref'
                    st.rerun()
            
            elif st.session_state.routine_step == 'study_pref':
                col1, col2, col3 = st.columns(3)
                with col1:
                    if st.button("🌅 Morning", key="morning_pref"):
                        st.session_state.project_data['routine']['study_preference'] = 'morning'
                        st.session_state.messages.append({"role": "user", "content": "I prefer morning study"})
                        st.session_state.messages.append({"role": "assistant", 
                                                         "content": "Great! Morning minds are fresh. What subjects are you studying?"})
                        st.session_state.step = 5
                        st.rerun()
                with col2:
                    if st.button("☀️ Afternoon", key="afternoon_pref"):
                        st.session_state.project_data['routine']['study_preference'] = 'afternoon'
                        st.session_state.messages.append({"role": "user", "content": "I prefer afternoon study"})
                        st.session_state.messages.append({"role": "assistant", 
                                                         "content": "Afternoon sessions work well! What subjects are you studying?"})
                        st.session_state.step = 5
                        st.rerun()
                with col3:
                    if st.button("🌙 Evening", key="evening_pref"):
                        st.session_state.project_data['routine']['study_preference'] = 'evening'
                        st.session_state.messages.append({"role": "user", "content": "I prefer evening study"})
                        st.session_state.messages.append({"role": "assistant", 
                                                         "content": "Evening study can be very productive! What subjects are you studying?"})
                        st.session_state.step = 5
                        st.rerun()
        
        # Step 5: Subjects
        elif st.session_state.step == 5:
            st.markdown("### Your Subjects")
            
            subjects_input = st.text_area(
                "Enter your subjects (comma-separated):",
                placeholder="e.g., Mathematics, Physics, Chemistry, Biology, English",
                height=80,
                key="subjects_input"
            )
            
            if st.button("Continue →", key="btn5", use_container_width=True):
                if subjects_input:
                    subjects = [s.strip() for s in subjects_input.split(',')]
                    st.session_state.project_data['subjects'] = subjects
                    st.session_state.messages.append({"role": "user", "content": subjects_input})
                    response = f"Perfect! I'll organize your {len(subjects)} subjects optimally. Now, let's choose a visual style for your planner!"
                    st.session_state.messages.append({"role": "assistant", "content": response})
                    st.session_state.step = 6
                    st.rerun()
                else:
                    st.error("Please enter at least one subject")
        
        # Step 6: Template Selection
        elif st.session_state.step == 6:
            st.markdown("### Choose Your Style")
            
            col1, col2, col3 = st.columns(3)
            
            with col1:
                st.markdown("""
                <div class="template-card template-simple">
                    <h3>📋 Simple</h3>
                    <p>Clean and straightforward</p>
                    <p style="font-size: 12px; color: var(--text-secondary);">
                        No distractions, just the essentials
                    </p>
                </div>
                """, unsafe_allow_html=True)
                if st.button("Choose Simple", key="simple_template"):
                    st.session_state.project_data['template'] = 'simple'
                    st.session_state.step = 7
                    st.rerun()
            
            with col2:
                st.markdown("""
                <div class="template-card template-minimal">
                    <h3>⚡ Minimal</h3>
                    <p>Modern and professional</p>
                    <p style="font-size: 12px; color: var(--text-secondary);">
                        Elegant design with subtle colors
                    </p>
                </div>
                """, unsafe_allow_html=True)
                if st.button("Choose Minimal", key="minimal_template"):
                    st.session_state.project_data['template'] = 'minimal'
                    st.session_state.step = 7
                    st.rerun()
            
            with col3:
                st.markdown("""
                <div class="template-card template-aesthetic">
                    <h3 style="color: white;">🎨 Aesthetic</h3>
                    <p style="color: white;">Colorful and motivating</p>
                    <p style="font-size: 12px; color: rgba(255,255,255,0.9);">
                        Beautiful gradients and emojis
                    </p>
                </div>
                """, unsafe_allow_html=True)
                if st.button("Choose Aesthetic", key="aesthetic_template"):
                    st.session_state.project_data['template'] = 'aesthetic'
                    st.session_state.step = 7
                    st.rerun()
        
        # Step 7: Generate Schedule
        elif st.session_state.step == 7:
            # Show loader
            show_loader("✨ Creating your personalized planner...", 3)
            
            # Generate schedule
            generator = ScheduleGenerator()
            schedule_df = generator.create_schedule(st.session_state.project_data)
            
            # Apply template styling
            styler = TemplateStyler()
            styled_df = styler.apply_style(schedule_df, st.session_state.project_data['template'])
            
            # Store generated plan
            st.session_state.project_data['generated_plan'] = styled_df
            
            # Show success message
            show_success_message("Your Planner is Ready!")
            
            # Display the schedule
            st.markdown("### 📊 Your Personalized Schedule")
            
            # Add custom styling based on template
            if st.session_state.project_data['template'] == 'aesthetic':
                st.markdown("""
                <style>
                    .dataframe {
                        background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
                        color: white;
                    }
                </style>
                """, unsafe_allow_html=True)
            
            st.dataframe(styled_df, use_container_width=True, height=400)
            
            # Export options
            st.markdown("### 💾 Download Your Planner")
            
            export_mgr = ExportManager()
            col1, col2, col3 = st.columns(3)
            
            with col1:
                pdf_bytes = export_mgr.to_pdf(styled_df, st.session_state.project_data)
                if pdf_bytes:
                    st.download_button(
                        label="📄 Download PDF",
                        data=pdf_bytes,
                        file_name=f"{st.session_state.project_data['folder_name']}_planner.pdf",
                        mime="application/pdf",
                        use_container_width=True
                    )
            
            with col2:
                excel_bytes = export_mgr.to_excel(styled_df, st.session_state.project_data)
                if excel_bytes:
                    st.download_button(
                        label="📊 Download Excel",
                        data=excel_bytes,
                        file_name=f"{st.session_state.project_data['folder_name']}_planner.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        use_container_width=True
                    )
            
            with col3:
                csv_bytes = export_mgr.to_csv(styled_df)
                if csv_bytes:
                    st.download_button(
                        label="📋 Download CSV",
                        data=csv_bytes,
                        file_name=f"{st.session_state.project_data['folder_name']}_planner.csv",
                        mime="text/csv",
                        use_container_width=True
                    )
            
            # Reset button
            st.markdown("---")
            if st.button("🔄 Create Another Planner", use_container_width=True):
                # Reset all session state
                for key in list(st.session_state.keys()):
                    del st.session_state[key]
                st.rerun()
    
    # Sidebar
    with st.sidebar:
        st.markdown("### 🎯 Planify Dashboard")
        
        # Progress metrics
        progress = (st.session_state.step / 7) * 100
        st.metric("Progress", f"{progress:.0f}%")
        st.progress(progress / 100)
        
        st.markdown("---")
        
        # Quick actions
        st.markdown("### ⚡ Quick Actions")
        
        if st.button("🔄 Start Over", use_container_width=True):
            for key in list(st.session_state.keys()):
                del st.session_state[key]
            st.rerun()
        
        if st.button("❓ Help", use_container_width=True):
            st.info("Need help? Follow the steps to create your personalized study plan!")
        
        st.markdown("---")
        
        # Tips
        st.markdown("### 💡 Tips")
        tips = [
            "Be specific about your routine for better results",
            "Choose a template that matches your style",
            "Update your plan weekly for best results",
            "Set realistic study hours",
            "Include breaks in your schedule"
        ]
        st.info(random.choice(tips))
        
        st.markdown("---")
        
        # About
        st.markdown("### ℹ️ About")
        st.markdown("""
        **Planify v1.0**  
        AI-Powered Study Planner
        
        Made with ❤️ for students
        """)

# ==================== RUN APPLICATION ====================
if __name__ == "__main__":
    main()